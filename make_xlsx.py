from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.comments import Comment

OUT = "weir_design.xlsx"
F = "Arial"
thin = Side(style="thin", color="A6A6A6")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
SUB_FILL = PatternFill("solid", fgColor="DDEBF7")
IN_FILL = PatternFill("solid", fgColor="FFFF00")
RES_FILL = PatternFill("solid", fgColor="E2EFDA")

wb = Workbook()


def title(ws, text, sub):
    ws["B1"] = text
    ws["B1"].font = Font(name=F, size=16, bold=True, color="1F4E78")
    ws["B2"] = sub
    ws["B2"].font = Font(name=F, size=10, italic=True, color="595959")


def header(ws, row, cols):
    for i, t in enumerate(cols):
        c = ws.cell(row=row, column=2 + i, value=t)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def section(ws, row, text, ncols):
    c = ws.cell(row=row, column=2, value=text)
    c.font = Font(name=F, bold=True, size=12, color="1F4E78")
    for i in range(ncols):
        ws.cell(row=row, column=2 + i).fill = SUB_FILL


def put(ws, row, values, fmts=None, inputs=(), results=()):
    for i, v in enumerate(values):
        col = 2 + i
        c = ws.cell(row=row, column=col, value=v)
        c.font = Font(name=F, color="0000FF" if col in inputs else "000000",
                      bold=col in results)
        c.border = BOX
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if col in inputs:
            c.fill = IN_FILL
        elif col in results:
            c.fill = RES_FILL
        if fmts and fmts[i]:
            c.number_format = fmts[i]
            c.alignment = Alignment(horizontal="center", vertical="center")


def widths(ws, w):
    for col, width in w.items():
        ws.column_dimensions[col].width = width
    ws.column_dimensions["A"].width = 2


# =================================================================== 1. ข้อมูล
ws = wb.active
ws.title = "1 ข้อมูล"
title(ws, "1) ข้อมูลออกแบบ — บ่อ 3 ช่อง: ปั๊มจุ่ม | โชว์น้ำใส | V-notch 90°",
      "แก้เฉพาะช่องสีเหลือง (ตัวเลขสีน้ำเงิน) แล้วชีต 2–5 จะคำนวณใหม่ให้เอง")
header(ws, 4, ["รายการ", "ค่า", "หน่วย", "ที่มา / หมายเหตุ"])
I = {}  # name -> absolute ref
rows = [
    ("qday", "น้ำเข้าต่อวัน", 1507, "m³/วัน", "จาก whiteboard"),
    ("qmin", "น้ำเข้าต่ำสุด", 50, "m³/h", "จากใบเสนอราคา (50–100 m³/h)"),
    ("qmax", "น้ำเข้าสูงสุด", 100, "m³/h", "จากใบเสนอราคา (50–100 m³/h)"),
    ("qr", "สูบไป Water recycle", 10, "m³/h", "จาก whiteboard"),
    ("b", "ความกว้างภายในบ่อ (B)", 1.20, "m", "เกณฑ์ ≥ 0.90 m และ h/B ≤ 0.2"),
    ("p", "ความสูงยอด V จากพื้นบ่อ (P)", 0.60, "m", "เกณฑ์ ≥ 0.45 m และ h/P ≤ 0.4"),
    ("notch", "ความลึกร่อง V", 0.30, "m", "เกณฑ์ ≥ hmax + 0.08 m"),
    ("lapp", "ความยาวช่วงตรงหน้าเวียร์", 2.20, "m", "เกณฑ์ ≥ 10 × hmax"),
    ("xtap", "ระยะจุดวัด h เหนือแผ่นเวียร์", 1.00, "m", "เกณฑ์ 4–5 × hmax"),
    ("raised", "ระดับผิวพื้นยก ช่อง 2", 0.90, "m", "วัดจากพื้นบ่อ"),
    ("sill", "ความสูงสันปลาย ช่อง 2", 0.02, "m", "สูงขึ้น = ชั้นน้ำหนาขึ้น ผิวน้ำเรียบขึ้น"),
]
r = 5
for key, label, val, unit, note in rows:
    put(ws, r, [label, val, unit, note], [None, "0.00" if val < 10 else "#,##0", None, None], inputs=(3,))
    I[key] = f"'1 ข้อมูล'!$C${r}"
    r += 1
put(ws, r, ["น้ำเข้าเฉลี่ย (= น้ำต่อวัน ÷ 24)", f"={I['qday']}/24", "m³/h", "คำนวณให้"], [None, "0.0", None, None])
I["qavg"] = f"'1 ข้อมูล'!$C${r}"
r += 2
section(ws, r, "ค่าคงที่ (ไม่ต้องแก้)", 4)
r += 1
consts = [
    ("g", "ความเร่งโน้มถ่วง (g)", 9.81, "m/s²", ""),
    ("ce", "สัมประสิทธิ์อัตราไหล V-notch 90° (Ce)", 0.578, "-", "ISO 1438 / USBR, สมการ Kindsvater–Shen"),
    ("kh", "ค่าปรับความสูงน้ำ V-notch 90° (kh)", 0.00085, "m", "ISO 1438 / USBR"),
    ("cf", "สัมประสิทธิ์สันปลายตรง (Francis)", 1.84, "-", "Q = 1.84 · B · H^1.5"),
]
for key, label, val, unit, note in consts:
    put(ws, r, [label, val, unit, note], [None, "0.00000" if val < 0.01 else "0.000", None, None])
    I[key] = f"'1 ข้อมูล'!$C${r}"
    r += 1
r += 1
section(ws, r, "สัญลักษณ์สี", 4)
put(ws, r + 1, ["ช่องสีเหลือง ตัวเลขสีน้ำเงิน", "50", "", "ค่าที่แก้ได้"], inputs=(3,))
put(ws, r + 2, ["ช่องสีเขียว ตัวหนา", "16.1", "", "ผลลัพธ์สำคัญ"], results=(3,))
put(ws, r + 3, ["ตัวเลขสีดำ", "0.014", "", "สูตรคำนวณ ห้ามพิมพ์ทับ"])
widths(ws, {"B": 40, "C": 12, "D": 10, "E": 46})
ws.freeze_panes = "B5"

# =================================================================== 2. เวียร์
ws = wb.create_sheet("2 เวียร์")
title(ws, "2) เวียร์ V-notch 90° — หาความสูงน้ำ h จากอัตราการไหล",
      "สูตร ISO 1438:  Q = Ce × (8/15) × √(2g) × (h + kh)^2.5   →   h = (Q ÷ K)^0.4 − kh   โดย K = Ce × (8/15) × √(2g)")
section(ws, 4, "ขั้นที่ 1  หาค่าคงที่ K", 6)
put(ws, 5, ["K = Ce × (8/15) × √(2g)", f"={I['ce']}*(8/15)*SQRT(2*{I['g']})", "-", "≈ 1.365 สำหรับ V-notch 90°", "", ""],
    [None, "0.0000", None, None, None, None])
K = "'2 เวียร์'!$C$5"
section(ws, 7, "ขั้นที่ 2  แปลงอัตราการไหล แล้วหา h", 6)
header(ws, 8, ["กรณี", "Q ผ่านเวียร์\n(m³/h)", "Q\n(m³/s)\n= ÷3600", "h\n(m)\n= (Q÷K)^0.4 − kh", "h\n(cm)", "ระดับน้ำ ช่อง 3\n(m จากพื้น)\n= P + h"])
cases = [
    ("น้ำเข้าต่ำสุด ปั๊มเดิน", f"={I['qmin']}-{I['qr']}"),
    ("น้ำเข้าเฉลี่ย ปั๊มเดิน", f"={I['qavg']}-{I['qr']}"),
    ("น้ำเข้าสูงสุด ปั๊มเดิน", f"={I['qmax']}-{I['qr']}"),
    ("น้ำเข้าสูงสุด ปั๊มหยุด (ใช้ออกแบบ)", f"={I['qmax']}"),
]
for i, (name, q) in enumerate(cases):
    rr = 9 + i
    put(ws, rr, [name, q, f"=C{rr}/3600", f"=(D{rr}/{K})^0.4-{I['kh']}", f"=E{rr}*100", f"={I['p']}+E{rr}"],
        [None, "0.0", "0.00000", "0.0000", "0.0", "0.000"], results=(6,))
ws["B13"].font = Font(name=F, bold=True)
HMAX = "'2 เวียร์'!$E$12"
section(ws, 15, "ขั้นที่ 3  ตรวจสอบเกณฑ์ (ใช้ hmax = กรณีสุดท้าย)", 6)
header(ws, 16, ["รายการตรวจ", "เกณฑ์", "ค่าที่ได้", "ผล", "", ""])
checks = [
    ("h ต่ำสุด ≥ 0.05 m (ต่ำกว่านี้วัดไม่แม่น)", "≥ 0.05 m", "=E9", '=IF(D17>=0.05,"ผ่าน ✔","ไม่ผ่าน ✘")', "0.000"),
    ("hmax ≤ 0.38 m", "≤ 0.38 m", "=E12", '=IF(D18<=0.38,"ผ่าน ✔","ไม่ผ่าน ✘")', "0.000"),
    ("h / P", "≤ 0.4", f"=E12/{I['p']}", '=IF(D19<=0.4,"ผ่าน ✔","ไม่ผ่าน ✘")', "0.00"),
    ("h / B", "≤ 0.2", f"=E12/{I['b']}", '=IF(D20<=0.2,"ผ่าน ✔","ไม่ผ่าน ✘")', "0.00"),
    ("P", "≥ 0.45 m", f"={I['p']}", '=IF(D21>=0.45,"ผ่าน ✔","ไม่ผ่าน ✘")', "0.00"),
    ("B", "≥ 0.90 m", f"={I['b']}", '=IF(D22>=0.9,"ผ่าน ✔","ไม่ผ่าน ✘")', "0.00"),
    ("ร่อง V ลึกกว่า hmax อย่างน้อย (m)", "≥ 0.08 m", f"={I['notch']}-E12", '=IF(D23>=0.08,"ผ่าน ✔","ไม่ผ่าน ✘")', "0.000"),
    ("ช่วงตรงหน้าเวียร์ ÷ hmax", "≥ 10 เท่า", f"={I['lapp']}/E12", '=IF(D24>=10,"ผ่าน ✔","ไม่ผ่าน ✘")', "0.0"),
    ("ระยะจุดวัด h ÷ hmax", "4–5 เท่า", f"={I['xtap']}/E12", '=IF(AND(D25>=4,D25<=5),"ผ่าน ✔","ไม่ผ่าน ✘")', "0.0"),
]
for i, (name, crit, val, res, fmt) in enumerate(checks):
    rr = 17 + i
    put(ws, rr, [name, crit, val, res, "", ""], [None, None, fmt, None, None, None], results=(5,))
widths(ws, {"B": 40, "C": 14, "D": 14, "E": 18, "F": 10, "G": 18})
ws.row_dimensions[8].height = 58
ws.freeze_panes = "B4"

# =================================================================== 3. ช่องโชว์น้ำ
ws = wb.create_sheet("3 ช่องโชว์น้ำ")
title(ws, "3) ช่อง 2 โชว์น้ำใส — ชั้นน้ำบนพื้นยกหนาเท่าไร",
      "สันปลาย:  Q = 1.84 × B × H^1.5  →  H = (Q ÷ (1.84 × B))^(2/3)   |   ความลึก = สันปลาย + H   |   Froude = v ÷ √(g × ลึก)  (< 0.5 = ผิวน้ำเรียบ)")
header(ws, 4, ["กรณี", "Q ผ่านช่อง 2\n(m³/h)", "Q\n(m³/s)", "ม่านน้ำบนสัน H\n(cm)", "ชั้นน้ำบนพื้นยก\n(cm)\n= สัน + H",
               "ความเร็ว v\n(m/s)\n= Q ÷ (B × ลึก)", "Froude", "ผิวน้ำ"])
for i, (name, q) in enumerate(cases):
    rr = 5 + i
    put(ws, rr, [name, q, f"=C{rr}/3600",
                 f"=(D{rr}/({I['cf']}*{I['b']}))^(2/3)*100",
                 f"={I['sill']}*100+E{rr}",
                 f"=D{rr}/({I['b']}*F{rr}/100)",
                 f"=G{rr}/SQRT({I['g']}*F{rr}/100)",
                 f'=IF(H{rr}<0.5,"เรียบ ✔",IF(H{rr}<0.7,"พอใช้","เป็นระลอก ✘"))'],
        [None, "0.0", "0.00000", "0.0", "0.0", "0.00", "0.00", None], results=(6, 9))
r = 10
section(ws, r, "อ่านผลอย่างไร", 8)
notes = [
    "• ชั้นน้ำบนพื้นยก = ความหนาน้ำที่คนจะมองเห็น (ช่องสีเขียว)",
    "• Froude < 0.5 ผิวน้ำเรียบเหมือนกระจก เห็นพื้นชัด / 0.7–1.3 ผิวน้ำเป็นระลอก ดูไม่ใส",
    "• อยากให้น้ำบางลง: ลดสันปลาย (ชีต 1) หรือเพิ่มความกว้าง B — แต่ Froude จะสูงขึ้น ให้ดูคอลัมน์ 'ผิวน้ำ' ประกอบ",
]
for j, t in enumerate(notes):
    c = ws.cell(row=r + 1 + j, column=2, value=t)
    c.font = Font(name=F)
widths(ws, {"B": 36, "C": 13, "D": 11, "E": 15, "F": 16, "G": 17, "H": 10, "I": 13})
ws.row_dimensions[4].height = 58
ws.freeze_panes = "B5"

# =================================================================== 4. ระดับ
ws = wb.create_sheet("4 ระดับ")
title(ws, "4) ระดับต่าง ๆ ของบ่อ (พื้นบ่อ = ±0.000 m)", "คำนวณที่น้ำเข้าสูงสุด ปั๊มหยุด (กรณีน้ำสูงที่สุด)")
header(ws, 4, ["ระดับ", "m", "คิดจาก"])
lv = [
    ("สันกำแพง", "=C6+0.2", "น้ำช่อง 1 + เผื่อ 0.20 m"),
    ("น้ำ ช่อง 1 (โดยประมาณ)", f"=C8+'3 ช่องโชว์น้ำ'!F8/100+1.5*'3 ช่องโชว์น้ำ'!G8^2/(2*{I['g']})", "พื้นยก + ชั้นน้ำ + การสูญเสียทางเข้า 1.5·v²/2g"),
    ("สันปลาย ช่อง 2", f"=C8+{I['sill']}", "พื้นยก + สันปลาย"),
    ("ผิวพื้นยก ช่อง 2", f"={I['raised']}", "ชีต 1"),
    ("ขอบบนร่อง V", f"={I['p']}+{I['notch']}", "P + ความลึกร่อง"),
    ("น้ำ ช่อง 3 สูงสุด", "='2 เวียร์'!G12", "P + hmax"),
    ("ยอด V (h = 0)", f"={I['p']}", "P"),
    ("น้ำท้ายเวียร์ต้องไม่เกิน", f"={I['p']}-0.06", "ต่ำกว่ายอด V อย่างน้อย 6 cm"),
    ("พื้นบ่อ", 0, ""),
]
for i, (name, f, how) in enumerate(lv):
    put(ws, 5 + i, [name, f, how], [None, "+0.000;-0.000;±0.000", None])
r = 15
section(ws, r, "ตรวจสอบ", 3)
put(ws, r + 1, ["สันปลายช่อง 2 สูงกว่าน้ำช่อง 3 (m)", "=C7-C10", '=IF(C16>=0.1,"ผ่าน ✔ ม่านน้ำตกอิสระ","ไม่ผ่าน ✘ น้ำช่อง 3 จะย้อนขึ้นช่อง 2")'],
    [None, "0.000", None], results=(4,))
put(ws, r + 2, ["สันกำแพงสูงกว่าน้ำช่อง 1 (m)", "=C5-C6", '=IF(C17>=0.15,"ผ่าน ✔","ไม่ผ่าน ✘")'],
    [None, "0.000", None], results=(4,))
widths(ws, {"B": 36, "C": 12, "D": 46})
ws.freeze_panes = "B5"

# =================================================================== 5. ตาราง h-Q
ws = wb.create_sheet("5 ตาราง h-Q")
title(ws, "5) ตาราง h → Q สำหรับอ่านค่าหน้างาน", "วัด h จากยอด V (cm) แล้วอ่าน Q   |   Q = K × (h + kh)^2.5")
header(ws, 4, ["h (cm)", "Q (m³/h)", "Q (L/s)", "Q (m³/วัน)"])
for i, h in enumerate(range(5, 26)):
    rr = 5 + i
    put(ws, rr, [h, f"={K}*(B{rr}/100+{I['kh']})^2.5*3600", f"=C{rr}/3.6", f"=C{rr}*24"],
        ["0", "0.0", "0.00", "#,##0"])
widths(ws, {"B": 12, "C": 14, "D": 12, "E": 14})
ws.freeze_panes = "B5"

for s in wb.worksheets:
    s.sheet_view.showGridLines = False
    s.sheet_properties.pageSetUpPr.fitToPage = True

ws = wb["1 ข้อมูล"]
ws["C6"].comment = Comment("ค่าจากใบเสนอราคา: ระบบเดิน 50–100 m³/h", "Claude")
wb.save(OUT)
print("saved", OUT)
