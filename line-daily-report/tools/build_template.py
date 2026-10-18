"""สร้างแม่แบบฟอร์ม Daily Report สำหรับ Google Sheets จากไฟล์ Excel ต้นฉบับของบริษัท

ใช้: python3 build_template.py <ต้นฉบับ.xlsx> <ไฟล์ผลลัพธ์.xlsx>

- เลื่อนฟอร์มลง 1 แถว เพื่อเว้นแถวบนสุดไว้วางโลโก้บริษัท (Google Sheets ใส่รูปในหัวกระดาษไม่ได้)
- เปลี่ยน checkbox ของ Excel (Form Control) เป็นตัวอักษร ☐ เพราะ Google Sheets ไม่รองรับ
- ลบข้อมูลตัวอย่าง ให้เหลือแต่ฟอร์มเปล่า
- เปลี่ยนฟอนต์ Angsana New เป็น Sarabun (Google Fonts) ขนาดที่ดูเท่ากัน
"""
import copy
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.page import PageMargins

SRC, OUT = sys.argv[1], sys.argv[2]
LAST_ROW, LAST_COL = 64, 7
FONT = 'Sarabun'
SIZE_MAP = {14: 10, 18: 13}
BOX = '☐'

src_wb = openpyxl.load_workbook(SRC)
src = src_wb.active
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Daily Report'
ws.sheet_view.showGridLines = False

for col in 'ABCDEFG':
    ws.column_dimensions[col].width = src.column_dimensions[col].width
# คอลัมน์ % กว้างขึ้นเล็กน้อย (เอาจากคอลัมน์ F) ให้หัวตาราง "เปอร์เซนต์ความคืบหน้า" ไม่ตัดคำ
ws.column_dimensions['E'].width += 1.8
ws.column_dimensions['F'].width -= 1.8

# แถว 1 = โลโก้, ฟอร์มเดิมเริ่มแถว 2
ws.row_dimensions[1].height = 58
for r in range(1, LAST_ROW + 1):
    h = src.row_dimensions[r].height
    if h:
        ws.row_dimensions[r + 1].height = h
    for c in range(1, LAST_COL + 1):
        s = src.cell(r, c)
        d = ws.cell(r + 1, c)
        d.value = s.value
        if s.has_style:
            d.font = copy.copy(s.font)
            d.border = copy.copy(s.border)
            d.fill = copy.copy(s.fill)
            d.number_format = s.number_format
            d.alignment = copy.copy(s.alignment)

merges = [m for m in src.merged_cells.ranges if m.max_col <= LAST_COL and m.max_row <= LAST_ROW]


def T(ref):
    """ตำแหน่งในฟอร์มเดิม -> ตำแหน่งในแม่แบบ (เลื่อนลง 1 แถว)"""
    col = ''.join(ch for ch in ref if ch.isalpha())
    return f'{col}{int(ref[len(col):]) + 1}'


def unmerge(rng):
    global merges
    merges = [m for m in merges if str(m) != rng]


# ---- ปรับฟอร์ม (อ้างอิงตำแหน่งในฟอร์มเดิม) ----
unmerge('B10:C10')
unmerge('D49:G49')
summary_rows = range(39, 49)
plan_rows = range(55, 60)
merges = [m for m in merges if not (m.min_row in summary_rows and m.min_row == m.max_row)]
merges = [m for m in merges if not (m.min_row in plan_rows and m.min_col == 1 and m.min_row == m.max_row)]

extra_merges = ['D49:E49', 'F49:G49']
extra_merges += [f'A{r}:D{r}' for r in summary_rows] + [f'F{r}:G{r}' for r in summary_rows]
extra_merges += [f'A{r}:C{r}' for r in plan_rows]
extra_merges += ['F58:G58']

# ข้อมูลตัวอย่าง -> ว่าง
for ref in ['E3', 'E4', 'E5', 'E6', 'B4', 'B5', 'B6', 'A8', 'B11', 'B12', 'D15', 'B50', 'A51', 'A53',
            'B62', 'D62', 'B63', 'D63']:
    ws[T(ref)].value = None
for r in list(summary_rows) + list(plan_rows):
    for col in 'ABCDEFG':
        if r in plan_rows and col in 'DEFG':
            continue
        ws[T(f'{col}{r}')].value = None
for r in range(9, 14):
    ws[T(f'F{r}')].value = '-'
    ws[T(f'F{r}')].font = Font(name=FONT, size=10)
for r in range(54, 60):
    ws[T(f'F{r}')].value = '-' if r != 59 else None

# checkbox -> ☐
machines = ['รถเครน ขนาด .......... ตัน', 'รถเฮี๊ยบ ขนาด .......... ตัน', 'รถปูน ขนาด .......... ตัน', 'รถยนต์']
for i, text in enumerate(machines):
    ws[T(f'D{9 + i}')].value = f'   {BOX}   {text}'
    ws[T(f'D{54 + i}')].value = f'   {BOX}   {text}'
ws[T('D13')].value = f'   {BOX}   อื่นๆ โปรดระบุ'
# แผนพรุ่งนี้: เดิมมีแถวว่างคั่นก่อน "อื่นๆ" ย้ายขึ้นมาต่อจากรถยนต์ เหลือแถวว่างไว้ล่างสุด
ws[T('D58')].value = f'   {BOX}   อื่นๆ โปรดระบุ'
ws[T('D58')].font = copy.copy(ws[T('D57')].font)
ws[T('D59')].value = None
ws[T('E38')].value = 'เปอร์เซนต์\nความคืบหน้า'
ws[T('E38')].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws[T('A10')].value = f'{BOX} ฝนตก'
ws[T('B10')].value = f'{BOX} แดดจัด'
ws[T('C10')].value = f'{BOX} ปกติ'
for ref in ['A10', 'B10', 'C10']:
    ws[T(ref)].alignment = Alignment(horizontal='center', vertical='center')
# เดิม B10:C10 ผสานกันอยู่ พอแยกช่องแล้วไม่ต้องมีเส้นคั่นกลาง
for ref in ['B10', 'C10']:
    b = copy.copy(ws[T(ref)].border)
    ws[T(ref)].border = Border(left=Side(), right=b.right if ref == 'C10' else Side(), top=b.top, bottom=b.bottom)
ws[T('A10')].border = Border(left=ws[T('A10')].border.left, top=ws[T('A10')].border.top, bottom=ws[T('A10')].border.bottom)
ws[T('A15')].value = f' {BOX} มีผลกระทบต่อการทำงาน'
ws[T('A15')].alignment = Alignment(horizontal='left', vertical='center')
ws[T('A16')].value = f' {BOX} ไม่มีผลกระทบต่อการทำงาน'
ws[T('A16')].alignment = Alignment(horizontal='left', vertical='center')
ws[T('D49')].value = f'   {BOX} ไม่มีอุบัติเหตุเกิดขึ้นในโครงการ'
ws[T('F49')].value = f'   {BOX} มีอุบัติเหตุเกิดขึ้นในโครงการ'
for ref in ['D49', 'F49']:
    ws[T(ref)].alignment = Alignment(horizontal='left', vertical='center')
    ws[T(ref)].font = Font(name=FONT, size=10)

# กรอบภาพถ่าย: แถว 18-35 สูงเท่ากันทุกแถว (ความสูงรวมเท่าเดิม) ให้คำนวณตำแหน่งรูปได้ง่าย
photo_rows = range(18, 36)
total = sum(src.row_dimensions[r].height for r in photo_rows)
for r in photo_rows:
    ws.row_dimensions[r + 1].height = round(total / len(photo_rows), 2)

# ช่องลายเซ็น / วันที่: จัดกลางแนวตั้งด้านล่าง
for ref in ['B62', 'D62']:
    ws[T(ref)].alignment = Alignment(horizontal='center', vertical='bottom')

# รหัสฟอร์มท้ายกระดาษ (แทน footer &F ของ Excel)
ws[T('A64')].value = 'FM_Daily Report_R02'
ws[T('A64')].font = Font(name=FONT, size=7)

# ฟอนต์
for row in ws.iter_rows(min_row=1, max_row=LAST_ROW + 1, max_col=LAST_COL):
    for cell in row:
        f = cell.font
        if f is None:
            continue
        size = SIZE_MAP.get(int(f.sz or 14), f.sz)
        cell.font = Font(name=FONT, size=size, bold=f.b, italic=f.i, color=f.color, underline=f.u)
ws[T('A64')].font = Font(name=FONT, size=7)

# ผสานช่อง (ทำหลังแก้ค่าและฟอนต์ เพราะช่องที่ถูกผสานแก้ไม่ได้)
for m in merges:
    ws.merge_cells(start_row=m.min_row + 1, start_column=m.min_col, end_row=m.max_row + 1, end_column=m.max_col)
for rng in extra_merges:
    a, b = rng.split(':')
    ws.merge_cells(f'{T(a)}:{T(b)}')

# แถว 1 เว้นไว้วางโลโก้บริษัท (บอทใส่ตามบริษัทที่เลือกตอนออกรายงาน)

ws.print_area = f'A1:G{LAST_ROW + 1}'
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.orientation = 'portrait'
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.print_options.horizontalCentered = True
ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.3, bottom=0.3)
wb.save(OUT)
print('saved', OUT)
